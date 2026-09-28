"""Turns an uploaded file (CSV/TSV, Excel, JSON, JSON Lines, WhatsApp, Telegram, X API data) into Posts.
Format guide for users: docs/data-format.md."""
import csv
import itertools
import json
import logging
import re
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo
from engine.schema import Post
from engine.zones import guess_timezone
from scenario.adapters.io_archive import load_io_archive, parse_list_field
from scenario.adapters.ira import load_ira
from scenario.adapters.telegram import is_telegram_export, load_telegram
from scenario.adapters.whatsapp import load_whatsapp
from scenario.adapters.x_api import is_x_data, x_posts

log = logging.getLogger(__name__)

# Accepted column names per field, compared after lower-casing and turning spaces/dashes into "_"
ALIASES = {
    "post_id": ["post_id", "id", "tweet_id", "tweetid", "status_id", "message_id", "unique_id"],
    "account_id": ["account_id", "user_id", "userid", "author_id", "account", "user", "author", "username", "screen_name",
                   "handle", "user_id_str", "user_screen_name", "user_username", "author_username", "from", "sender",
                   "from_id", "channel", "page_name"],
    "username": ["username", "screen_name", "user_screen_name", "user_username", "author_username", "handle", "author",
                 "user", "account", "user_name", "name", "from"],
    "created_at": ["created_at", "timestamp", "time", "date", "datetime", "tweet_time", "posted_at", "publish_date",
                   "created", "created_time", "published", "published_at", "date_time", "post_date", "sent_at"],
    "text": ["text", "content", "message", "post", "tweet", "tweet_text", "body"],
    "urls": ["urls", "url", "links", "link"],
    "hashtags": ["hashtags", "hashtag", "tags"],
    "reply_to": ["reply_to", "in_reply_to", "in_reply_to_id", "in_reply_to_status_id", "in_reply_to_tweetid", "parent_id"],
    "repost_of": ["repost_of", "retweet_of", "retweet_id", "retweet_tweetid", "retweeted_status_id", "shared_post_id"],
    "account_created_at": ["account_created_at", "account_creation_date", "user_created_at"],
    "platform": ["platform", "source", "network", "app", "channel_type"],
    "city": ["city", "town", "location", "district", "place", "user_location", "user_reported_location"],
    "language": ["language", "lang", "tweet_language"],
    # shown in the Posts view when present
    "display_name": ["display_name", "user_display_name", "name", "user_name", "author_name", "user_name_display"],
    "quote_of": ["quote_of", "quoted_tweet_tweetid", "quoted_tweet_id", "quoted_status_id", "quoted_status_id_str",
                 "quoted_post_id"],
    "conversation_id": ["conversation_id", "thread_id"],
    "reply_to_user": ["reply_to_user", "in_reply_to_screen_name", "in_reply_to_username"],
    "followers": ["followers", "follower_count", "followers_count", "user_followers_count"],
    "verified": ["verified", "user_verified", "is_verified"],
    "likes": ["likes", "like_count", "favorite_count", "favourite_count", "reactions", "reaction_count"],
    "reposts": ["reposts", "retweet_count", "repost_count", "shares", "share_count"],
    "replies": ["replies", "reply_count", "comment_count", "comments"],
    "quotes": ["quotes", "quote_count"],
    "views": ["views", "view_count", "impression_count", "impressions"],
}
METRICS = ("likes", "reposts", "replies", "quotes", "views")
# Fields the column-matching step asks about; the others are always matched by name
FORM_FIELDS = {"account_id", "created_at", "text", "username", "post_id", "platform", "city", "language", "urls",
               "hashtags", "reply_to", "repost_of", "account_created_at"}
REQUIRED = {"account_id": "who posted it", "created_at": "when it was posted", "text": "what it says"}
JSON_SUFFIXES = (".json", ".jsonl", ".ndjson")
# ponytail: 05/09/2020 is read month-first (US exports); day-first is used only when the first number is over 12
TIME_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y",
                "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
                "%a %b %d %H:%M:%S %z %Y", "%d %b %Y %H:%M", "%d %B %Y %H:%M", "%b %d, %Y %I:%M %p"]
URL_RE = re.compile(r"https?://\S+")
# \w alone stops at vowel signs and other combining marks, which cut "#कोरोना" to "#क"
HASHTAG_RE = re.compile(r"#[\w\u0300-\u036F\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED\u0900-\u0DFF\u0E00-\u0E7F]+")
# Non-Latin scripts name the language well enough for triage (Cyrillic is reported as Russian, Arabic script with
# Urdu letters as Urdu). Checked in this order: Japanese kana before Chinese characters, which Japanese also uses.
SCRIPTS = [
    ("hi", r"[\u0900-\u097F]"), ("bn", r"[\u0980-\u09FF]"), ("pa", r"[\u0A00-\u0A7F]"), ("gu", r"[\u0A80-\u0AFF]"),
    ("or", r"[\u0B00-\u0B7F]"), ("ta", r"[\u0B80-\u0BFF]"), ("te", r"[\u0C00-\u0C7F]"), ("kn", r"[\u0C80-\u0CFF]"),
    ("ml", r"[\u0D00-\u0D7F]"), ("ur", r"[\u0679\u0688\u0691\u06BA\u06BE\u06C1\u06D2]"), ("ar", r"[\u0600-\u06FF]"),
    ("ru", r"[\u0400-\u04FF]"), ("el", r"[\u0370-\u03FF]"), ("he", r"[\u0590-\u05FF]"), ("th", r"[\u0E00-\u0E7F]"),
    ("ja", r"[\u3040-\u30FF]"), ("ko", r"[\uAC00-\uD7AF]"), ("zh", r"[\u4E00-\u9FFF]"),
]
SCRIPT_RES = [(code, re.compile(pattern)) for code, pattern in SCRIPTS]
# Hindi in Roman script: needs at least one unambiguous word, so Spanish or Catalan "se", "ke" don't count
HINGLISH_STRONG = {"hai", "hain", "nahi", "nahin", "kya", "aaj", "bhai", "mein", "raha", "rahi", "rahe", "karo", "kar",
                   "aur", "bhi", "yeh", "woh", "abhi", "tum", "hum", "sabhi", "gaya", "gayi", "wala", "wali", "diya"}
HINGLISH_WEAK = {"ke", "ki", "ko", "se", "ye", "pe", "sab", "log", "kal"}
STOPWORDS = {  # the most frequent words of each language; the language with most hits (at least 2) wins
    "en": {"the", "and", "is", "are", "to", "of", "in", "for", "you", "this", "that", "with", "on", "it", "was", "be",
           "not", "have", "will", "they", "we", "at"},
    "es": {"el", "los", "las", "que", "y", "en", "por", "para", "con", "una", "es", "del", "se", "no", "lo", "como"},
    "ca": {"els", "les", "que", "i", "per", "amb", "una", "és", "del", "es", "no", "al", "més", "dels", "ho", "això"},
    "fr": {"le", "les", "et", "est", "des", "un", "une", "pour", "dans", "que", "qui", "pas", "du", "au", "sur"},
    "de": {"der", "die", "das", "und", "ist", "nicht", "mit", "ein", "eine", "zu", "den", "von", "auf", "sich"},
    "pt": {"os", "que", "e", "em", "para", "com", "um", "uma", "não", "do", "da", "dos", "se", "no"},
}
LANGUAGE_NAMES = dict(zip(
    "english hindi russian ukrainian belarusian bulgarian serbian macedonian croatian bosnian slovenian polish czech "
    "slovak hungarian romanian arabic urdu farsi persian pushto pashto turkish azerbaijani kazakh uzbek spanish "
    "catalan french german dutch italian portuguese swedish norwegian danish finnish estonian latvian lithuanian "
    "greek hebrew chinese japanese korean vietnamese indonesian malay thai bengali tamil telugu marathi gujarati "
    "punjabi kannada malayalam".split(),
    "en hi ru uk be bg sr mk hr bs sl pl cs sk hu ro ar ur fa fa ps ps tr az kk uz es ca fr de nl it pt sv no da fi et "
    "lv lt el he zh ja ko vi id ms th bn ta te mr gu pa kn ml".split()))
UNKNOWN_LANGUAGE = {"", "und", "qme", "zxx", "qht", "qam", "qct", "qst", "unknown", "lang", "none",  # X's "no language"
                    "language undefined", "undefined"}


def detect_language(text: str) -> str | None:
    """Script first (Devanagari → hi, Cyrillic → ru, …); for Latin script, Hinglish words, then common words."""
    text = text or ""
    letters = re.sub(r"https?://\S+|[@#]\w+", "", text)
    for code, pattern in SCRIPT_RES:
        if len(pattern.findall(letters)) >= 3:
            return code
    words = re.findall(r"[a-z\u00E0-\u00FF]+", letters.lower())
    strong = sum(w in HINGLISH_STRONG for w in words)
    if strong and strong + sum(w in HINGLISH_WEAK for w in words) >= 2:
        return "hinglish"
    hits = {code: sum(w in stop for w in words) for code, stop in STOPWORDS.items()}
    best = max(hits, key=hits.get)
    return best if hits[best] >= 2 else ("en" if hits["en"] else None)


def normalize_language(value: str | None) -> str | None:
    """A language column as a code: 'Russian' → 'ru', 'EN' → 'en', X's 'und' → None."""
    v = str(value or "").strip().lower()
    return None if v in UNKNOWN_LANGUAGE else LANGUAGE_NAMES.get(v, v)


def parse_time(value, naive_zone=None) -> int | None:
    """Unix seconds from epoch seconds/milliseconds or a date string. Times without a zone are read in naive_zone
    (the dataset's local clock), UTC if none is given."""
    v = str(value if value is not None else "").strip()
    if not v:
        return None
    if re.fullmatch(r"\d+(\.\d+)?", v):
        t = float(v)
        return int(t / 1000 if t > 1e11 else t)  # 13-digit values are milliseconds
    try:
        dt = datetime.fromisoformat(v.replace("Z", "+00:00"))
    except ValueError:
        for fmt in TIME_FORMATS:
            try:
                dt = datetime.strptime(v, fmt)
                break
            except ValueError:
                continue
        else:
            return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=naive_zone or timezone.utc)
    return int(dt.timestamp())


ZONE_MARK_RE = re.compile(r"(\dZ|[+-]\d\d:?\d\d|(?i:UTC|GMT))(\s+\d{4})?\s*$")  # also "…14:05:12 +0000 2026"


def has_zone(value) -> bool:
    """Whether a time value pins down its moment: Unix time, or a date string ending in Z, +05:30, UTC…"""
    v = str(value if value is not None else "").strip()
    return bool(re.fullmatch(r"\d+(\.\d+)?", v) or ZONE_MARK_RE.search(v))


def as_list(value) -> list[str]:
    if isinstance(value, list):  # [{"tag": "x"}] from nested JSON exports, or plain strings
        return [str(x.get("tag") or x.get("text") or x.get("expanded_url") or x.get("url") or "") if isinstance(x, dict)
                else str(x) for x in value if x]
    return parse_list_field(str(value or ""))


def as_int(value) -> int | None:
    try:
        return int(float(str(value).replace(",", "")))
    except (TypeError, ValueError):
        return None


def column_key(name: str) -> str:
    return re.sub(r"[\s\-.]+", "_", str(name).strip().lower())  # "Tweet ID" = "tweet-id" = "tweet.id" = tweet_id


class MissingColumns(ValueError):
    """The file is a table, but which column is the account, time or text could not be worked out from the names."""


def looks_like_time(value) -> bool:
    t = parse_time(value)
    return t is not None and 946684800 <= t <= 4102444800  # 2000-2100: a count or an ID is not a date


def suggest_mapping(columns: list[str], sample: list[dict]) -> dict[str, str | None]:
    """Which column holds each field: by name first (ALIASES), then by the values for the three that are needed:
    the time is the column of dates, the text the column of long strings, the account a short value that repeats."""
    keys = {c: column_key(c) for c in columns}
    mapping = {field: next((c for name in names for c in columns if keys[c] == name), None) for field, names in ALIASES.items()}
    values = {c: [str(r.get(c) if r.get(c) is not None else "").strip() for r in sample] for c in columns}
    filled = lambda c: [v for v in values[c] if v]
    free = lambda: [c for c in columns if c not in set(mapping.values())]

    if not mapping["created_at"]:
        dated = [(sum(map(looks_like_time, filled(c))) / max(1, len(filled(c))), c) for c in free() if filled(c)]
        share, best = max(dated, default=(0, None))
        mapping["created_at"] = best if share >= 0.8 else None
    if not mapping["text"]:
        lengths = [(sum(map(len, filled(c))) / max(1, len(filled(c))), c) for c in free() if filled(c)]
        avg, best = max(lengths, default=(0, None))
        mapping["text"] = best if avg >= 15 else None
    if not mapping["account_id"]:
        hints = ("from", "sender", "name", "handle", "user", "author", "by", "poster", "owner", "channel", "page")
        candidates = []
        for c in free():
            vals = filled(c)
            if len(vals) < 0.9 * len(sample) or not vals or sum(map(len, vals)) / len(vals) > 40:
                continue
            distinct = len(set(vals)) / len(vals)
            if distinct < 1 or len(vals) < 5:  # someone posts more than once
                candidates.append((any(h in keys[c] for h in hints), -distinct, c))
        mapping["account_id"] = max(candidates)[2] if candidates else None
    if mapping["username"] == mapping["account_id"]:
        mapping["username"] = None
    return mapping


def resolve_columns(columns: list[str], mapping: dict | None = None) -> dict[str, str | None]:
    """The column for each field: from a mapping the user confirmed, else by name (ALIASES)."""
    keys = {c: column_key(c) for c in columns}
    by_name = {field: next((c for name in names for c in columns if keys[c] == name), None) for field, names in ALIASES.items()}
    if not mapping:
        return by_name
    return {field: (mapping.get(field) if mapping.get(field) in columns else None) if field in FORM_FIELDS else by_name[field]
            for field in ALIASES}


def load_rows(rows, columns: list[str], mapping: dict | None = None, naive_zone=None) -> list[Post]:
    """Map rows (dicts) onto Posts. Times without a zone are read in naive_zone (UTC if None)."""
    col = resolve_columns(columns, mapping)

    missing = [f for f in REQUIRED if not col[f]]
    if missing:
        need = "; ".join(f"{REQUIRED[f]} (e.g. {' / '.join(ALIASES[f][:4])})" for f in missing)
        hint = (" The engine finds campaigns from who posted what and when, so a text-only dataset can't be analysed"
                " for coordination." if col["text"] else "")
        raise MissingColumns(f"Missing columns for {need}. Found: {', '.join(columns[:30])}.{hint}")

    posts, skipped = [], 0
    for i, row in enumerate(rows):
        def get(field):
            return row.get(col[field]) if col[field] else None

        account = str(get("account_id") or "").strip()
        created_at = parse_time(get("created_at"), naive_zone)
        if not account or created_at is None:
            skipped += 1
            continue
        text = str(get("text") or "")
        posts.append(Post(
            post_id=str(get("post_id") or f"row{i + 1}"),
            account_id=account,
            username=str(get("username") or account),
            created_at=created_at,
            text=text,
            repost_of=str(get("repost_of") or "").strip() or None,
            reply_to=str(get("reply_to") or "").strip() or None,
            # no column (or empty): take links and hashtags from the text
            urls=as_list(get("urls")) or URL_RE.findall(text),
            hashtags=as_list(get("hashtags")) or HASHTAG_RE.findall(text),
            account_created_at=parse_time(get("account_created_at"), naive_zone),
            platform=str(get("platform") or "").strip().lower() or None,
            city=str(get("city") or "").strip() or None,
            language=str(get("language") or "").strip().lower() or None,
            display_name=str(get("display_name") or "").strip() or None,
            quote_of=str(get("quote_of") or "").strip() or None,
            conversation_id=str(get("conversation_id") or "").strip() or None,
            reply_to_user=str(get("reply_to_user") or "").strip().lstrip("@") or None,
            followers=as_int(get("followers")),
            verified=str(get("verified")).strip().lower() in ("true", "1", "yes") if get("verified") not in (None, "") else None,
            metrics={m: v for m in METRICS if (v := as_int(get(m))) is not None},
        ))

    if not posts:
        raise ValueError(f"No usable rows: all {skipped} rows lack an account or a readable time "
                         f"(time column '{col['created_at']}'). See docs/data-format.md for accepted time formats.")
    if skipped:
        log.warning("Skipped %d rows without an account or a readable time", skipped)
    return posts


def load_posts(path: str | Path, limit: int | None = None, mapping: dict | None = None) -> list[Post]:
    """Any supported file as Posts. `mapping` (field -> column) overrides the automatic column matching."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Dataset file not found: {path}")
    posts = None if mapping else _load_known_format(path, limit)
    if posts is not None:
        return enrich(posts)
    with open_table(path, limit) as (columns, rows):
        posts = enrich(load_rows(rows, columns, mapping))
    # "22/09/2020 10:00" means 10:00 on the poster's clock: if the posts are Indian (IST), read such times again as IST
    zone = guess_timezone(posts)
    if zone != "UTC":
        with open_table(path, 50) as (columns, rows):
            time_col = resolve_columns(columns, mapping)["created_at"]
            naive = any(r.get(time_col) and not has_zone(r.get(time_col)) for r in rows)
        if naive:
            with open_table(path, limit) as (columns, rows):
                posts = enrich(load_rows(rows, columns, mapping, naive_zone=ZoneInfo(zone)))
    return posts


def enrich(posts: list[Post]) -> list[Post]:
    for p in posts:  # every source gets language, links and hashtags, from its columns or from the text
        p.language = normalize_language(p.language) or detect_language(p.text)
        p.urls = p.urls or URL_RE.findall(p.text)
        p.hashtags = p.hashtags or HASHTAG_RE.findall(p.text)
    return posts


def read_json(path: Path):
    """A JSON document, or JSON Lines (one object per line, as X streams and twarc write them)."""
    text = path.read_text(encoding="utf-8-sig")
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return [json.loads(line) for line in text.splitlines() if line.strip()]


def _load_known_format(path: Path, limit: int | None) -> list[Post] | None:
    """Exports with their own adapter: WhatsApp, Telegram, X API data, the IO and IRA research archives.
    None for an ordinary table."""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        posts = load_whatsapp(path, limit=limit)
        if not posts:
            raise ValueError("No messages found. Expected a WhatsApp chat export (lines like '22/09/20, 10:02 - Name: message').")
        return posts
    if suffix in JSON_SUFFIXES:
        data = read_json(path)
        if is_telegram_export(data):
            return load_telegram(path, limit=limit)
        if is_x_data(data):
            return x_posts(data)[:limit] if limit else x_posts(data)
        return None
    if suffix in (".csv", ".tsv"):
        with open_table(path, 0) as (columns, _):
            names = {column_key(c) for c in columns}
        # Known research archives need their own handling (retweet fields, rounded IDs)
        if {"tweetid", "tweet_text"} <= names:
            return load_io_archive(path, limit=limit)
        if "publish_date" in names and ("author" in names or "tweet_id" in names):
            return load_ira(path, limit=limit)
    return None


def flatten(row: dict, prefix: str = "") -> dict:
    """Nested JSON as flat columns: {"user": {"screen_name": "x"}} -> {"user.screen_name": "x"}."""
    flat = {}
    for k, v in row.items():
        if isinstance(v, dict) and v:
            flat.update(flatten(v, f"{prefix}{k}."))
        else:
            flat[f"{prefix}{k}"] = v
    return flat


@contextmanager
def open_table(path: Path, limit: int | None = None):
    """(columns, rows) of a CSV/TSV, Excel sheet, JSON array or JSON Lines file. Rows are read lazily for CSV."""
    suffix = path.suffix.lower()
    if suffix in JSON_SUFFIXES:
        data = read_json(path)
        if isinstance(data, dict):  # {"posts": [...]} or similar: the first list of objects inside
            data = next((v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)), [data])
        if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
            raise ValueError("Expected JSON objects, one per post (an array, or one per line)")
        rows = [flatten(r) for r in (data[:limit] if limit is not None else data)]
        yield list(dict.fromkeys(k for r in rows[:1000] for k in r)), rows
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            cells = book.worksheets[0].iter_rows(values_only=True)
            header = [str(h if h is not None else f"column {i + 1}") for i, h in enumerate(next(cells, []))]
            cell = lambda v: v.isoformat() if hasattr(v, "isoformat") else v  # Excel dates arrive as datetime
            rows = ({h: cell(v) for h, v in zip(header, r)} for r in cells)
            yield header, itertools.islice(rows, limit) if limit is not None else rows
        finally:
            book.close()
    else:
        with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
            try:  # comma, semicolon (Excel in many locales) or tab
                dialect = csv.Sniffer().sniff(f.read(64 * 1024), delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
            f.seek(0)
            reader = csv.DictReader(f, dialect=dialect)
            yield reader.fieldnames or [], itertools.islice(reader, limit) if limit is not None else reader


def preview(path: str | Path, mapping: dict | None = None) -> dict:
    """First rows and the suggested column for each field, for the column-matching step of an upload."""
    with open_table(Path(path), 200) as (columns, rows):
        sample = list(rows)
    short = lambda v: (s if len(s := str(v if v is not None else "")) <= 80 else s[:77] + "…")
    return {
        "columns": columns,
        "rows": [{c: short(r.get(c)) for c in columns} for r in sample[:5]],
        "suggested": mapping or suggest_mapping(columns, sample),
    }
